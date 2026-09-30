"""Email → job pipeline: poll, match sender rules, classify, check links, create/dedupe
jobs, schedule processing, and hand due jobs to the generation / portal pipelines.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Literal
from urllib.parse import parse_qsl, urlencode, urlparse
from zoneinfo import ZoneInfo

import pymupdf
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.boards.linkedin import AlertPosting, alert_postings, is_linkedin_alert, title_skipped
from app.boards.linkedin import canonical_url as linkedin_url
from app.boards.service import prepare_linkedin_job
from app.core.config import get_config
from app.email.links import LinkResolver, check_link, extract_links, unwrap
from app.email.parse import ParsedEmail, auth_passed, forwarded_original_sender, parse_mime
from app.email.providers import Candidate, EmailError, MailProvider, first_sync_since
from app.generation.level import gate_and_queue
from app.knowledge.skills import normalize
from app.llm.errors import LLMError
from app.llm.gateway import LLMGateway
from app.models import EmailAccount, Job, JobFile, Portal, SenderRule, SourceEmail
from app.models.enums import (
    ApplyMode,
    ConnectionStatus,
    EmailClassification,
    JobFileKind,
    JobSource,
    JobStatus,
    LLMTask,
    NotificationLevel,
)
from app.schemas.settings import AppSettings
from app.services.notifications import notify
from app.services.settings_service import get_app_settings

log = logging.getLogger(__name__)
BODY_PROMPT_CHARS = 8000


# -- rule matching --------------------------------------------------------------------------


def sender_matches(sender: str, pattern: str) -> bool:
    sender, p = sender.lower().strip(), pattern.lower().strip()
    if not sender or not p:
        return False
    if "@" in p and not p.startswith("@"):
        return sender == p
    domain = p.lstrip("@")
    host = sender.rsplit("@", 1)[-1]
    return host == domain or host.endswith("." + domain)


def candidate_rule(c: Candidate, rules: list[SenderRule], own: set[str]) -> SenderRule | None:
    """Cheap pre-filter on metadata only: a rule sender, or (for forwarding rules) one of
    the user's own addresses. Nothing else is ever downloaded."""
    for r in rules:
        if r.enabled and sender_matches(c.sender, r.sender_match):
            return r
    if c.sender in own:
        return next((r for r in rules if r.enabled and r.match_forwarded), None)
    return None


@dataclass
class SenderDecision:
    rule: SenderRule | None
    effective_sender: str
    forwarded: bool
    reason: str


def resolve_sender(parsed: ParsedEmail, rules: list[SenderRule], own: set[str]) -> SenderDecision:
    """Direct mail must match a rule and must not fail DMARC/DKIM. Forwarded mail counts
    only if the forwarder is one of the user's own addresses, the forward itself passes
    authentication, and the original sender inside matches a forwarding-enabled rule."""
    auth = auth_passed(parsed)
    for r in rules:
        if r.enabled and sender_matches(parsed.sender, r.sender_match):
            if auth is False:
                return SenderDecision(
                    None,
                    parsed.sender,
                    False,
                    "sender failed DMARC/DKIM authentication (possible spoof)",
                )
            return SenderDecision(r, parsed.sender, False, "sender allowlisted")
    if parsed.sender in own:
        original = forwarded_original_sender(parsed)
        if original:
            if auth is False:
                return SenderDecision(None, original, True, "forward failed authentication")
            for r in rules:
                if r.enabled and r.match_forwarded and sender_matches(original, r.sender_match):
                    return SenderDecision(r, original, True, f"forwarded from {original}")
    return SenderDecision(None, parsed.sender, False, "sender not on the allowlist")


# -- classification -----------------------------------------------------------------------


class EmailClassificationOut(BaseModel):
    kind: Literal["job_notice", "deadline_extension", "reminder", "result", "general_notice"]
    company: str | None = None
    role: str | None = None
    deadline: str | None = Field(default=None, description="ISO 8601, e.g. 2026-09-30T23:59")
    apply_link: int | None = Field(default=None, description="number of the apply/job link")
    ctc: str | None = None
    location: str | None = None


def classify(
    gateway: LLMGateway, parsed: ParsedEmail, links: list[str], job_id: int | None = None
) -> EmailClassificationOut:
    result = gateway.structured(
        LLMTask.EMAIL_CLASSIFIER,
        "email_classify",
        {
            "sender": parsed.sender,
            "subject": parsed.subject,
            "body": parsed.body_text[:BODY_PROMPT_CHARS],
            "links": [f"{i + 1}. {u}" for i, u in enumerate(links)],
        },
        EmailClassificationOut,
        job_id=job_id,
    )
    return result.value


def parse_deadline(value: str | None, tz_name: str) -> datetime | None:
    if not value:
        return None
    try:
        d = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError:
        return None
    if d.tzinfo is None:
        d = d.replace(tzinfo=ZoneInfo(tz_name))
    return d.astimezone(UTC)


# -- jobs -----------------------------------------------------------------------------------

_ID_PARAMS = re.compile(r"^(id|job|jobid|job_id|opening|opportunity|drive|posting)", re.I)


def dedupe_key(
    portal_id: int | None, url: str | None, company: str | None, role: str | None
) -> str:
    """Same opportunity from both inboxes / reminders → same key."""
    if url:
        u = urlparse(url)
        params = sorted((k, v) for k, v in parse_qsl(u.query) if _ID_PARAMS.match(k))
        basis = (
            f"url:{portal_id}:{(u.hostname or '').lower()}{u.path.rstrip('/')}?{urlencode(params)}"
        )
    else:
        basis = f"text:{portal_id}:{normalize(company or '')}:{normalize(role or '')}"
    return hashlib.sha256(basis.encode()).hexdigest()[:40]


def schedule(
    job: Job, rule: SenderRule | None, received: datetime, global_delay: int
) -> str | None:
    """Set scheduled_at from the delay (job > rule > global). Returns a note if the
    deadline forces immediate processing."""
    delay = (
        job.delay_minutes
        if job.delay_minutes is not None
        else (rule.delay_minutes if rule and rule.delay_minutes is not None else global_delay)
    )
    job.scheduled_at = received + timedelta(minutes=delay)
    deadline = (
        job.deadline.replace(tzinfo=UTC)
        if job.deadline and job.deadline.tzinfo is None
        else job.deadline
    )
    if deadline and deadline < job.scheduled_at:
        job.scheduled_at = datetime.now(UTC)
        return (
            f"deadline {deadline:%d %b %H:%M UTC} is sooner than the {delay}-minute delay; "
            "processing immediately"
        )
    return None


@dataclass
class PollReport:
    account: str
    candidates: int = 0
    downloaded: int = 0
    jobs_created: int = 0
    duplicates: int = 0
    suspicious: int = 0
    ignored: int = 0
    updated: int = 0
    errors: list[str] = field(default_factory=list)


def _own_addresses(db: Session) -> set[str]:
    return {a.lower() for a in db.scalars(select(EmailAccount.address))}


def _save_attachments(db: Session, job: Job, parsed: ParsedEmail) -> None:
    root = Path(get_config().data_dir) / "jobs" / str(job.id)
    root.mkdir(parents=True, exist_ok=True)
    for i, a in enumerate(parsed.attachments):
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", a.filename)[:120] or f"attachment{i}"
        path = root / f"email_{i}_{safe}"
        path.write_bytes(a.data)
        db.add(
            JobFile(
                job_id=job.id,
                kind=JobFileKind.EMAIL_ATTACHMENT,
                path=str(path.relative_to(get_config().data_dir)),
                original_name=a.filename,
                mime_type=a.mime_type,
                sha256=hashlib.sha256(a.data).hexdigest(),
            )
        )


def ingest_message(
    db: Session,
    gateway: LLMGateway,
    account: EmailAccount,
    parsed: ParsedEmail,
    rules: list[SenderRule],
    own: set[str],
    resolver: LinkResolver | None,
    report: PollReport,
) -> SourceEmail | None:
    """Three phases so no DB transaction is ever open during network calls:
    1) read-only checks, 2) LLM classification + link resolution, 3) one write txn."""
    # -- 1. read-only checks --------------------------------------------------------------
    if db.scalar(
        select(SourceEmail.id).where(
            SourceEmail.account_id == account.id, SourceEmail.message_id == parsed.message_id
        )
    ):
        return None
    decision = resolve_sender(parsed, rules, own)
    if decision.rule is None:
        report.ignored += 1
        return None  # not stored: only allowlisted mail is kept
    rule = decision.rule
    twin = None
    if parsed.internet_message_id:
        twin = db.scalar(
            select(SourceEmail).where(
                SourceEmail.internet_message_id == parsed.internet_message_id,
                SourceEmail.account_id != account.id,
            )
        )
    portal = db.get(Portal, rule.portal_id) if rule.portal_id else None
    allowed = list(portal.allowed_domains) if portal else []
    portal_id = portal.id if portal else None
    settings = get_app_settings(db)

    def new_source() -> SourceEmail:
        return SourceEmail(
            account_id=account.id,
            message_id=parsed.message_id,
            internet_message_id=parsed.internet_message_id,
            sender=parsed.sender,
            original_sender=decision.effective_sender if decision.forwarded else None,
            subject=parsed.subject,
            received_at=parsed.received_at,
            body_html=parsed.html,
            body_text=parsed.body_text,
            headers=parsed.headers,
            sender_rule_id=rule.id,
        )

    if twin is not None:  # the same message arrived in the other inbox: same job
        src = new_source()
        src.classification, src.job_id = twin.classification, twin.job_id
        src.extracted_link, src.resolved_link, src.link_safe = (
            twin.extracted_link,
            twin.resolved_link,
            twin.link_safe,
        )
        db.add(src)
        report.duplicates += 1
        db.commit()
        return src
    db.commit()  # end the read transaction before any network call

    postings = (
        alert_postings(parsed.html, parsed.text)
        if is_linkedin_alert(decision.effective_sender or parsed.sender, parsed.html, parsed.text)
        else []
    )
    if postings:
        return _ingest_linkedin_alert(db, new_source(), rule, parsed, postings, settings, report)

    # -- 2. network: classify, pick and check the link -----------------------------------
    candidates = extract_links(parsed.html, parsed.text)
    link_urls = [unwrap(c.url) for c in candidates][:15]
    try:
        cls: EmailClassificationOut | None = classify(gateway, parsed, link_urls)
    except LLMError as exc:
        cls = None
        report.errors.append(f"classification failed for '{parsed.subject[:60]}': {exc}"[:300])
    chosen: str | None = None
    safety = None
    if cls is not None and cls.kind == "job_notice":
        if cls.apply_link and 1 <= cls.apply_link <= len(link_urls):
            chosen = candidates[cls.apply_link - 1].url
        elif candidates:
            chosen = candidates[0].url
        if chosen:
            safety = check_link(
                chosen, sender_allowed=True, allowed_domains=allowed, resolver=resolver
            )

    # -- 3. write ---------------------------------------------------------------------------
    src = new_source()
    db.add(src)
    if cls is None:
        src.classification = EmailClassification.UNCLASSIFIED
        db.commit()
        return src
    src.classification = EmailClassification(cls.kind)
    deadline = parse_deadline(cls.deadline, settings.timezone)

    if cls.kind != "job_notice":
        existing = _find_existing_job(db, portal_id, link_urls, cls)
        if existing is not None:
            src.job_id = existing.id
            if cls.kind == "deadline_extension" and deadline:
                old = existing.deadline
                existing.deadline = deadline
                report.updated += 1
                notify(
                    db,
                    NotificationLevel.INFO,
                    "job.deadline_extended",
                    f"Deadline extended: {existing.company} – {existing.role}",
                    f"{old:%d %b %H:%M} → {deadline:%d %b %H:%M} UTC" if old else None,
                    job_id=existing.id,
                )
        db.commit()
        return src

    src.extracted_link = chosen
    if safety is not None:
        src.resolved_link, src.link_safe, src.link_check_reason = (
            safety.final_url,
            safety.safe,
            safety.reason,
        )
    key = dedupe_key(
        portal_id, safety.final_url if safety and safety.safe else None, cls.company, cls.role
    )
    job = db.scalar(select(Job).where(Job.dedupe_key == key))
    if job is not None:
        src.job_id = job.id
        report.duplicates += 1
        if deadline and (job.deadline is None or deadline > job.deadline.replace(tzinfo=UTC)):
            job.deadline = deadline
        db.commit()
        return src

    job = Job(
        source=JobSource.EMAIL,
        status=JobStatus.DETECTED,
        company=cls.company,
        role=cls.role,
        portal_id=portal_id,
        apply_mode=rule.apply_mode,
        dedupe_key=key,
        deadline=deadline,
        ctc=cls.ctc,
        location=cls.location,
        detected_at=datetime.now(UTC),
    )
    db.add(job)
    db.flush()
    src.job_id = job.id
    report.jobs_created += 1
    if safety is None and rule.apply_mode is ApplyMode.READ_EMAIL_THEN_APPLY:
        # No link to open at all: the email itself is the JD. Nothing unsafe was opened,
        # so generate from the email; applying is then up to the user.
        job.status_reason = (
            "No apply link in the email: the resume is generated from the "
            "email; apply yourself after reviewing it."
        )
        note = schedule(job, rule, parsed.received_at, settings.global_delay_minutes)
        _save_attachments(db, job, parsed)
        db.commit()
        if note:
            notify(
                db,
                NotificationLevel.WARNING,
                "job.urgent",
                f"Urgent: {cls.company or 'job'} – {cls.role or ''}",
                note,
                job_id=job.id,
            )
        return src
    if safety is None or not safety.safe:
        job.status = JobStatus.SUSPICIOUS
        job.status_reason = (
            "No apply link found in the email"
            if safety is None
            else f"Link not opened: {safety.reason}"
        )
        report.suspicious += 1
        _save_attachments(db, job, parsed)
        db.commit()
        notify(
            db,
            NotificationLevel.WARNING,
            "job.suspicious",
            f"Suspicious link in email: {parsed.subject[:80]}",
            f"{job.status_reason}. Link: {chosen or '(none)'}. Nothing was opened.",
            job_id=job.id,
        )
        return src
    job.apply_url = safety.final_url
    note = schedule(job, rule, parsed.received_at, settings.global_delay_minutes)
    _save_attachments(db, job, parsed)
    db.commit()
    if note:
        notify(
            db,
            NotificationLevel.WARNING,
            "job.urgent",
            f"Urgent: {cls.company or 'job'} – {cls.role or ''}",
            note,
            job_id=job.id,
        )
    return src


def _ingest_linkedin_alert(
    db: Session,
    src: SourceEmail,
    rule: SenderRule,
    parsed: ParsedEmail,
    postings: list[AlertPosting],
    settings: AppSettings,
    report: PollReport,
) -> SourceEmail:
    """A LinkedIn alert lists several postings: one job each (deduped by LinkedIn job id).
    Nothing is opened here; the posting is read when the job's delay has passed."""
    src.classification = EmailClassification.JOB_NOTICE
    src.extracted_link = linkedin_url(postings[0].job_id)
    db.add(src)
    first: Job | None = None
    for p in postings:
        if title_skipped(p.title, settings.skip_title_words):
            continue  # e.g. senior roles: no resume, no quota spent
        key = f"linkedin:{p.job_id}"
        job = db.scalar(select(Job).where(Job.dedupe_key == key))
        if job is None:
            job = Job(
                source=JobSource.EMAIL,
                status=JobStatus.DETECTED,
                role=p.title or None,
                apply_mode=rule.apply_mode,
                dedupe_key=key,
                portal_job_ref=key,
                apply_url=linkedin_url(p.job_id),
                # fallback JD if the public posting can't be read (replaced when it can)
                jd_text=(p.snippet or p.title or None),
                detected_at=datetime.now(UTC),
            )
            db.add(job)
            db.flush()
            schedule(job, rule, parsed.received_at, settings.global_delay_minutes)
            report.jobs_created += 1
        else:
            report.duplicates += 1
        first = first or job
    src.job_id = first.id if first else None
    db.commit()
    return src


def _find_existing_job(
    db: Session, pid: int | None, links: list[str], cls: EmailClassificationOut
) -> Job | None:
    for url in links:
        job = db.scalar(select(Job).where(Job.dedupe_key == dedupe_key(pid, url, None, None)))
        if job is not None:
            return job
    if cls.company:
        return db.scalar(
            select(Job).where(Job.dedupe_key == dedupe_key(pid, None, cls.company, cls.role))
        )
    return None


def poll_account(
    db: Session,
    gateway: LLMGateway,
    account: EmailAccount,
    provider: MailProvider,
    resolver: LinkResolver | None,
) -> PollReport:
    report = PollReport(account=account.address)
    rules = list(db.scalars(select(SenderRule).where(SenderRule.email_account_id == account.id)))
    own = _own_addresses(db)
    try:
        candidates, cursor = provider.list_candidates(
            account, first_sync_since(), account.sync_cursor
        )
        report.candidates = len(candidates)
        for c in candidates:
            if candidate_rule(c, rules, own) is None:
                continue
            if db.scalar(
                select(SourceEmail.id).where(
                    SourceEmail.account_id == account.id, SourceEmail.message_id == c.provider_id
                )
            ):
                continue
            parsed = parse_mime(
                provider.fetch_raw(account, c.provider_id), provider_id=c.provider_id
            )
            report.downloaded += 1
            ingest_message(db, gateway, account, parsed, rules, own, resolver, report)
        account.sync_cursor = cursor
        account.status, account.last_error = ConnectionStatus.CONNECTED, None
    except (EmailError, Exception) as exc:  # noqa: BLE001 — surface any provider failure
        account.status = ConnectionStatus.ERROR
        account.last_error = str(exc)[:500]
        report.errors.append(str(exc)[:300])
        log.warning("poll %s failed: %s", account.address, exc)
    account.last_sync_at = datetime.now(UTC)
    db.commit()
    return report


# -- due jobs ------------------------------------------------------------------------------


def email_jd_text(db: Session, job: Job) -> str:
    """JD text for 'read email then apply': subject + body (tables flattened) + attachments."""
    parts: list[str] = []
    for src in db.scalars(
        select(SourceEmail).where(SourceEmail.job_id == job.id).order_by(SourceEmail.received_at)
    ):
        parts.append(f"Subject: {src.subject}\n\n{src.body_text or ''}")
    for f in db.scalars(
        select(JobFile).where(
            JobFile.job_id == job.id, JobFile.kind == JobFileKind.EMAIL_ATTACHMENT
        )
    ):
        content = attachment_text(Path(get_config().data_dir) / f.path)
        parts.append(f"Attachment {f.original_name}:\n{content}")
    return "\n\n".join(p for p in parts if p.strip())[:60_000]


def attachment_text(path: Path) -> str:
    try:
        if path.suffix.lower() == ".pdf" or path.name.lower().endswith(".pdf"):
            with pymupdf.open(path) as doc:
                return "\n".join(p.get_text("text", sort=True) for p in doc)
        if path.name.lower().endswith(".docx"):
            import docx  # python-docx

            d = docx.Document(io.BytesIO(path.read_bytes()))
            rows = [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
            return "\n".join([p.text for p in d.paragraphs] + rows)
        if path.name.lower().endswith(".txt"):
            return path.read_text(errors="replace")
    except Exception as exc:  # noqa: BLE001 — a broken attachment shouldn't stop the job
        return f"(could not read attachment: {type(exc).__name__})"
    return ""


def process_due_jobs(
    db: Session,
    queue_generation: Callable[..., object],
    queue_portal: Callable[..., object] | None,
    gateway: LLMGateway | None = None,
) -> int:
    """Hand scheduled email jobs to the next stage once their delay has passed."""
    now = datetime.now(UTC)
    due = list(
        db.scalars(
            select(Job).where(
                Job.source.in_([JobSource.EMAIL, JobSource.SITE]),
                Job.status == JobStatus.DETECTED,
                Job.scheduled_at.is_not(None),
                Job.scheduled_at <= now,
            )
        )
    )
    for job in due:
        job.scheduled_at = None  # claimed
        if (job.portal_job_ref or "").startswith("linkedin:"):
            job.jd_text = job.jd_text or email_jd_text(db, job)  # fallback if unreadable
            db.commit()
            prepare_linkedin_job(db, job.id)  # full JD + the company's own apply page
            gate_and_queue(db, job, queue_generation, gateway)  # fresher/experienced (AI)
            continue
        if job.source is JobSource.SITE and queue_portal is not None:
            db.commit()  # stage 2: open the posting, read the JD (+PDFs), then generate
            queue_portal(db, job, then_generate=True, merge_email=False)
            continue
        if job.apply_mode is ApplyMode.READ_EMAIL_THEN_APPLY or queue_portal is None:
            job.jd_text = email_jd_text(db, job)
            db.commit()
            portal = db.get(Portal, job.portal_id) if job.portal_id else None
            portal_readable = bool(job.apply_url and portal and portal.credentials)
            if (
                queue_portal is not None
                and job.apply_mode is ApplyMode.READ_EMAIL_THEN_APPLY
                and portal_readable
            ):
                queue_portal(db, job, then_generate=True, merge_email=True)
            else:
                gate_and_queue(db, job, queue_generation, gateway)
        else:  # direct link: the JD comes from the portal page
            db.commit()
            queue_portal(db, job, then_generate=True, merge_email=False)
    return len(due)
